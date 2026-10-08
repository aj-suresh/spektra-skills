#!/usr/bin/env python3
"""Read a .docx draft and insert hyperlinks into a copy of it.

  docx_links.py extract DRAFT.docx [--json OUT.json]
      Numbered paragraphs ([P12] text), heading levels, and the links already
      in the draft. --json also writes the structured version that
      candidates.py reads.

  docx_links.py apply DRAFT.docx LINKS.json OUT.docx
      LINKS.json is a list of
        {"paragraph": 12, "anchor": "exact text in that paragraph",
         "url": "https://...", "occurrence": 1}
      Each anchor is wrapped in a real Word hyperlink, keeping the run
      formatting (bold, italics, font) and adding the usual blue underline.
      The original file is never modified. Links that can't be placed
      cleanly are reported, not forced.

  docx_links.py verify OUT.docx LINKS.json
      Confirms every planned link is present with the right URL.

Paragraph numbers count every paragraph in the body, including table cells,
in document order, the same way in all three commands.

Needs python-docx (run through scripts/run.sh, which installs it).
"""
import argparse
import copy
import json
import re
import sys

try:
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.exit("python-docx is missing. Run this through scripts/run.sh, or: pip install python-docx")

W_P, W_R, W_T, W_HYPER = qn("w:p"), qn("w:r"), qn("w:t"), qn("w:hyperlink")
R_ID = qn("r:id")


def paragraphs(document):
    """Every w:p in the body, document order (tables included)."""
    return [p for p in document.element.body.iter(W_P)]


def style_of(p_el, document):
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        return ""
    ps = ppr.find(qn("w:pStyle"))
    if ps is None:
        return ""
    sid = ps.get(qn("w:val"))
    try:
        return document.styles.get_by_id(sid, 1).name or sid  # WD_STYLE_TYPE.PARAGRAPH == 1
    except Exception:
        return sid


def heading_level(style, p_el):
    m = re.match(r"(?:heading|título|titre|überschrift)\s*(\d)", style or "", re.I)
    if m:
        return int(m.group(1))
    if (style or "").lower() == "title":
        return 0
    ppr = p_el.find(qn("w:pPr"))
    if ppr is not None:
        ol = ppr.find(qn("w:outlineLvl"))
        if ol is not None:
            return int(ol.get(qn("w:val"))) + 1
    return None


def run_text(r):
    out = []
    for c in r:
        if c.tag == W_T:
            out.append(c.text or "")
        elif c.tag == qn("w:tab"):
            out.append("\t")
        elif c.tag in (qn("w:br"), qn("w:cr")):
            out.append("\n")
    return "".join(out)


def para_text(p_el):
    out = []
    for c in p_el:
        if c.tag == W_R:
            out.append(run_text(c))
        elif c.tag == W_HYPER or c.tag in (qn("w:ins"), qn("w:smartTag"), qn("w:fldSimple")):
            for r in c.iter(W_R):
                out.append(run_text(r))
    return "".join(out)


def para_links(p_el, part):
    links = []
    for h in p_el.iter(W_HYPER):
        rid = h.get(R_ID)
        url = part.rels[rid].target_ref if rid and rid in part.rels else ("#" + (h.get(qn("w:anchor")) or ""))
        links.append({"text": "".join(run_text(r) for r in h.iter(W_R)), "url": url})
    # HYPERLINK field codes (older Word docs)
    for instr in p_el.iter(qn("w:instrText")):
        m = re.search(r'HYPERLINK\s+"([^"]+)"', instr.text or "")
        if m:
            links.append({"text": "(field)", "url": m.group(1)})
    return links


def cmd_extract(args):
    d = docx.Document(args.docx)
    part = d.part
    rows, words = [], 0
    for i, p in enumerate(paragraphs(d)):
        text = para_text(p)
        if not text.strip():
            continue
        style = style_of(p, d)
        lvl = heading_level(style, p)
        links = para_links(p, part)
        wc = len(text.split())
        words += wc if lvl is None else 0
        in_table = any(a.tag == qn("w:tc") for a in p.iterancestors())
        is_list = "list" in (style or "").lower() or p.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None
        rows.append({"i": i, "style": style, "heading": lvl, "table": in_table, "list": is_list,
                     "text": text, "words": wc, "links": links})
    for r in rows:
        tags = ["P%d" % r["i"]]
        if r["heading"] is not None:
            tags.append("TITLE" if r["heading"] == 0 else "H%d" % r["heading"])
        if r["table"]:
            tags.append("table")
        if r["list"]:
            tags.append("list")
        print("[%s] %s" % (" · ".join(tags), r["text"]))
        for l in r["links"]:
            print("      ↳ existing link: %r -> %s" % (l["text"], l["url"]))
    print("\n-- %d paragraphs, ~%d body words, %d existing links" %
          (len(rows), words, sum(len(r["links"]) for r in rows)), file=sys.stderr)
    if args.json:
        with open(args.json, "w") as f:
            json.dump({"source": args.docx, "body_words": words, "paragraphs": rows}, f, indent=1)


def split_run(r, k):
    """Split simple run r at text offset k -> (left, right); both keep rPr."""
    t = r.find(W_T)
    text = t.text or ""
    left = copy.deepcopy(r)
    left.find(W_T).text = text[:k]
    t.text = text[k:]
    for el in (left.find(W_T), t):
        el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.addprevious(left)
    return left, r


def simple(r):
    """A run we can split safely: only rPr plus w:t children. Merge multiple w:t."""
    kids = [c for c in r if c.tag != qn("w:rPr")]
    if not kids or any(c.tag != W_T for c in kids):
        return False
    if len(kids) > 1:
        kids[0].text = "".join(c.text or "" for c in kids)
        for c in kids[1:]:
            r.remove(c)
    return True


def style_link_run(r, document):
    rpr = r.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        r.insert(0, rpr)
    try:
        has_style = any(s.style_id == "Hyperlink" for s in document.styles)
    except Exception:
        has_style = False
    if has_style and rpr.find(qn("w:rStyle")) is None:
        rs = OxmlElement("w:rStyle")
        rs.set(qn("w:val"), "Hyperlink")
        rpr.insert(0, rs)
    for tag, attrs in (("w:color", {"w:val": "1155CC"}), ("w:u", {"w:val": "single"})):
        old = rpr.find(qn(tag))
        if old is not None:
            rpr.remove(old)
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        rpr.append(el)


def insert_link(document, p_el, anchor, url, occurrence=1):
    runs = [c for c in p_el if c.tag == W_R]
    spans, pos = [], 0
    for r in runs:
        t = run_text(r)
        spans.append((r, pos, pos + len(t)))
        pos += len(t)
    full = "".join(run_text(r) for r in runs)
    idx, start = -1, 0
    for _ in range(max(1, occurrence)):
        idx = full.find(anchor, start)
        if idx < 0:
            break
        start = idx + 1
    if idx < 0:
        if anchor in para_text(p_el):
            return "anchor is inside an existing link or tracked change"
        return "anchor text not found in paragraph"
    end = idx + len(anchor)
    hit = [(r, a, b) for (r, a, b) in spans if b > idx and a < end]
    if not all(simple(r) for r, _, _ in hit):
        return "anchor spans a tab/line break/image; place manually"
    first, a0, _ = hit[0]
    if idx > a0:
        _, first = split_run(first, idx - a0)
        hit[0] = (first, idx, hit[0][2])
    last, a1, b1 = hit[-1]
    if end < b1:
        last, _ = split_run(last, end - a1)
        hit[-1] = (last, a1, end)
    covered = [r for r, _, _ in hit]
    rid = document.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(R_ID, rid)
    covered[0].addprevious(h)
    for r in covered:
        style_link_run(r, document)
        h.append(r)
    return None


def cmd_apply(args):
    d = docx.Document(args.docx)
    plist = paragraphs(d)
    plan = json.load(open(args.links))
    if isinstance(plan, dict):
        plan = plan.get("links", [])
    ok, failed = 0, []
    for l in plan:
        i = l["paragraph"]
        if i >= len(plist):
            failed.append((l, "no paragraph %d" % i))
            continue
        err = insert_link(d, plist[i], l["anchor"], l["url"], l.get("occurrence", 1))
        if err:
            failed.append((l, err))
        else:
            ok += 1
    d.save(args.out)
    print("inserted %d/%d links -> %s" % (ok, len(plan), args.out))
    for l, err in failed:
        print("  NOT INSERTED [P%s] %r -> %s : %s" % (l["paragraph"], l["anchor"], l["url"], err))
    sys.exit(0 if not failed else 3)


def cmd_verify(args):
    d = docx.Document(args.docx)
    plist = paragraphs(d)
    plan = json.load(open(args.links))
    if isinstance(plan, dict):
        plan = plan.get("links", [])
    missing = 0
    for l in plan:
        found = any(x["text"] == l["anchor"] and x["url"] == l["url"]
                    for x in para_links(plist[l["paragraph"]], d.part)) if l["paragraph"] < len(plist) else False
        print("%s [P%d] %r -> %s" % ("OK  " if found else "MISS", l["paragraph"], l["anchor"], l["url"]))
        missing += not found
    total = sum(len(para_links(p, d.part)) for p in plist)
    print("-- %d/%d planned links present; %d hyperlinks in document overall" % (len(plan) - missing, len(plan), total))
    sys.exit(0 if not missing else 3)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("extract")
    e.add_argument("docx")
    e.add_argument("--json")
    a = sub.add_parser("apply")
    a.add_argument("docx")
    a.add_argument("links")
    a.add_argument("out")
    v = sub.add_parser("verify")
    v.add_argument("docx")
    v.add_argument("links")
    args = ap.parse_args()
    {"extract": cmd_extract, "apply": cmd_apply, "verify": cmd_verify}[args.cmd](args)


if __name__ == "__main__":
    main()
